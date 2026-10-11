import os
import re
import json
import io
import hmac
import threading
import time
import urllib.request
import urllib.error
from datetime import datetime

from flask import Flask, render_template, jsonify, send_file, request, Response
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from invoice_parser import (
    parse_invoice, compute_interventions, layout_text, parse_ddmmyyyy, ParseError,
)

app = Flask(__name__)

# ---------------------------------------------------------------- configurare
GOOGLE_CREDENTIALS_JSON = os.environ.get("GOOGLE_CREDENTIALS_JSON", "")
FOLDER_ID = os.environ.get("DRIVE_FOLDER_ID", "")
POLL_INTERVAL = int(os.environ.get("POLL_INTERVAL", "60"))

# AI folosit DOAR ca rezerva, cand o factura nu are sablonul SimNettoyage.
# Implicit: Groq (gratuit), model openai/gpt-oss-120b. Orice API compatibil OpenAI merge.
GROQ_API_KEY = os.environ.get("GROQ_API_KEY", "").strip()
LLM_API_KEY = os.environ.get("LLM_API_KEY", "").strip() or GROQ_API_KEY
LLM_BASE_URL = os.environ.get("LLM_BASE_URL", "https://api.groq.com/openai/v1").rstrip("/")
LLM_MODEL = os.environ.get("LLM_MODEL", "openai/gpt-oss-120b")

# Optional: parola pentru dashboard (HTTP Basic Auth, orice utilizator)
APP_PASSWORD = os.environ.get("APP_PASSWORD", "")

state = {
    "files": {},
    "interventii": [],
    "last_scan": None,
    "monitoring": False,
    "log": [],
}
state_lock = threading.Lock()
scan_lock = threading.Lock()


def add_log(msg):
    ts = datetime.now().strftime("%H:%M:%S")
    with state_lock:
        state["log"].append(f"[{ts}] {msg}")
        if len(state["log"]) > 200:
            state["log"] = state["log"][-200:]


# ---------------------------------------------------------------- autentificare
@app.before_request
def require_password():
    if not APP_PASSWORD or request.path == "/healthz":
        return None
    auth = request.authorization
    if auth and auth.password and hmac.compare_digest(auth.password, APP_PASSWORD):
        return None
    return Response("Autentificare necesara", 401,
                    {"WWW-Authenticate": 'Basic realm="Facturi"'})


# ---------------------------------------------------------------- Google Drive
def get_drive_service():
    if not GOOGLE_CREDENTIALS_JSON:
        raise ValueError("GOOGLE_CREDENTIALS_JSON not set")
    creds = service_account.Credentials.from_service_account_info(
        json.loads(GOOGLE_CREDENTIALS_JSON),
        scopes=["https://www.googleapis.com/auth/drive.readonly"],
    )
    return build("drive", "v3", credentials=creds, cache_discovery=False)


def list_pdf_files(folder_id):
    service = get_drive_service()
    all_files = []

    def search_folder(fid):
        page_token = None
        while True:  # paginare: altfel se pierd fisierele peste primele 100
            results = service.files().list(
                q=f"'{fid}' in parents and trashed=false",
                fields="nextPageToken, files(id, name, mimeType, modifiedTime, size)",
                pageSize=1000,
                pageToken=page_token,
                supportsAllDrives=True,
                includeItemsFromAllDrives=True,
            ).execute()
            for f in results.get("files", []):
                if f["mimeType"] == "application/pdf":
                    all_files.append(f)
                elif f["mimeType"] == "application/vnd.google-apps.folder":
                    search_folder(f["id"])
            page_token = results.get("nextPageToken")
            if not page_token:
                break

    search_folder(folder_id)
    return all_files


def download_pdf(file_id):
    service = get_drive_service()
    req = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    buf = io.BytesIO()
    downloader = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- rezerva AI
def llm_extract(pdf_bytes, max_retries=3):
    """
    Pentru facturi care NU au sablonul SimNettoyage. AI-ul doar copiaza datele
    brute (client, randuri, date); calculul interventiilor se face tot in cod.
    """
    text = layout_text(pdf_bytes)
    text = re.sub(r"[ \t]{3,}", " | ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)[:12000]

    prompt = (
        "Extract the data from this French cleaning-company invoice. "
        "The issuer is SimNettoyage — ignore its address and bank details.\n"
        "Return ONLY a JSON object, no markdown, with this shape:\n"
        '{"client": "name right under \'Facture pour\' (first line only, no address)",'
        ' "numar": "invoice number after \'Facture #\'",'
        ' "data_emitere": "DD/MM/YYYY from \'Soumis le\'",'
        ' "randuri": [{"nom": "property name", "type": "service type", '
        '"dates": ["dates exactly as written"], "qty": number, "pu": unit price number, '
        '"total": line total number}]}\n'
        "Use the 'Détail des prestations' table if present (one entry per property row; "
        "a row can have several dates). Skip empty €0.00 rows. Copy numbers, do not compute.\n\n"
        "INVOICE:\n" + text
    )
    body = json.dumps({
        "model": LLM_MODEL,
        "messages": [{"role": "user", "content": prompt}],
        "temperature": 0,
        "max_tokens": 4000,
    }).encode()

    for attempt in range(1, max_retries + 1):
        req = urllib.request.Request(
            f"{LLM_BASE_URL}/chat/completions", data=body,
            headers={"Content-Type": "application/json",
                     "Authorization": f"Bearer {LLM_API_KEY}",
                     "User-Agent": "Cleaning-App/2.0"},
        )
        try:
            with urllib.request.urlopen(req, timeout=90) as resp:
                data = json.loads(resp.read())
            raw = (data["choices"][0]["message"].get("content") or "").strip()
            s, e = raw.find("{"), raw.rfind("}") + 1
            if s < 0 or e <= s:
                raise ValueError("raspuns AI fara JSON")
            return json.loads(raw[s:e])
        except urllib.error.HTTPError as e:
            if e.code == 429 and attempt < max_retries:
                wait = int(e.headers.get("retry-after") or 20)
                add_log(f"Limita AI atinsa, astept {wait}s (incercarea {attempt}/{max_retries})")
                time.sleep(min(wait, 60))
                continue
            raise
        except (ValueError, KeyError) as e:
            if attempt < max_retries:
                add_log(f"Raspuns AI invalid ({e}), reincerc...")
                continue
            raise
    raise RuntimeError("AI indisponibil")


def parse_with_ai(pdf_bytes):
    data = llm_extract(pdf_bytes)
    header = {
        "client": (data.get("client") or "").strip() or None,
        "numar": (data.get("numar") or "").strip() or None,
        "data_emitere": parse_ddmmyyyy(data.get("data_emitere") or ""),
    }
    rows = []
    for r in data.get("randuri") or []:
        def num(v):
            try:
                return float(str(v).replace("€", "").replace(",", ".").strip()) if v not in (None, "") else None
            except ValueError:
                return None
        rows.append({"nom": (r.get("nom") or "").strip(), "type": r.get("type"),
                     "dates_raw": [str(d) for d in (r.get("dates") or [])],
                     "qty": num(r.get("qty")), "pu": num(r.get("pu")), "total": num(r.get("total"))})
    rows = [r for r in rows if r["nom"] or r["total"]]
    interventii, warnings = compute_interventions(rows, header)
    total = round(sum(r["total"] or 0 for r in rows), 2)
    return {
        "metoda": "ai",
        "antet": {**header, "data_emitere": header["data_emitere"].strftime("%d/%m/%Y")
                  if header["data_emitere"] else None},
        "randuri": rows,
        "totaluri": {"subtotal": total, "ajustari": 0.0, "total": total},
        "interventii": interventii,
        "avertismente": ["Citita cu AI (sablon necunoscut) — verificati manual"] + warnings,
    }


# ---------------------------------------------------------------- procesare
def _dup_key(result):
    a, t = result["antet"], result["totaluri"]
    return ((a.get("numar") or "").upper(), (a.get("client") or "").casefold(), t.get("total"))


def process_file(file_info):
    file_id, name = file_info["id"], file_info["name"]
    with state_lock:
        entry = state["files"].setdefault(file_id, {})
        entry.update({
            "id": file_id, "name": name, "status": "processing",
            "modifiedTime": file_info.get("modifiedTime", entry.get("modifiedTime")),
            "error": None, "avertismente": [],
        })
        entry.setdefault("interventii", [])
    used_ai = False
    try:
        pdf_bytes = download_pdf(file_id)
        try:
            result = parse_invoice(pdf_bytes)
        except ParseError as pe:
            if not LLM_API_KEY:
                raise ValueError(f"{pe} (nicio cheie AI pentru rezerva)")
            add_log(f"{name}: {pe} — incerc cu AI ({LLM_MODEL})")
            used_ai = True
            result = parse_with_ai(pdf_bytes)

        interventii = result["interventii"]
        for iv in interventii:
            iv["_source"] = name
            iv["_file_id"] = file_id
        a, t = result["antet"], result["totaluri"]
        warnings = list(result["avertismente"])

        with state_lock:
            others = [f for fid, f in state["files"].items()
                      if fid != file_id and f.get("status") in ("done", "warning")]
            dup = next((f for f in others if f.get("_dup_key") == _dup_key(result)), None)
            same_nr = [f for f in others if a.get("numar") and f.get("numar") == a.get("numar")]

            entry = state["files"][file_id]
            entry.update({
                "numar": a.get("numar"), "client": a.get("client"),
                "data_emitere": a.get("data_emitere"), "total": t.get("total"),
                "metoda": result["metoda"], "_dup_key": _dup_key(result),
            })
            if dup:
                entry["status"] = "duplicate"
                entry["avertismente"] = [f"Duplicat al fisierului {dup['name']} — ignorat"]
                entry["interventii"] = []
                state["interventii"] = [iv for iv in state["interventii"] if iv.get("_file_id") != file_id]
            else:
                for f in same_nr:
                    warnings.append(f"Numarul {a['numar']} e folosit si de {f['name']} (client {f.get('client')})")
                entry["status"] = "warning" if warnings else "done"
                entry["avertismente"] = warnings
                entry["interventii"] = interventii
                state["interventii"] = [
                    iv for iv in state["interventii"] if iv.get("_file_id") != file_id
                ] + interventii

        if dup:
            add_log(f"= {name}: duplicat al {dup['name']}, ignorat")
        else:
            add_log(f"✓ {name}: {a.get('numar')} · {a.get('client')} · "
                    f"{len(interventii)} interventii · {t.get('total')} €"
                    + (f" · {len(warnings)} avertismente" if warnings else ""))
            for w in warnings:
                add_log(f"  ⚠ {w}")
    except Exception as e:
        with state_lock:
            state["files"][file_id]["status"] = "error"
            state["files"][file_id]["error"] = str(e)
        add_log(f"✗ Eroare la {name}: {e}")
    return used_ai


def process_queue_background(new_files):
    add_log(f"Procesez {len(new_files)} fisiere...")
    for i, f in enumerate(new_files):
        used_ai = process_file(f)
        if used_ai and i < len(new_files) - 1:
            time.sleep(3)  # menajam limita gratuita a AI-ului
    add_log("Toate fisierele procesate.")


def scan_folder():
    if not FOLDER_ID:
        add_log("DRIVE_FOLDER_ID nu e setat")
        return
    if not scan_lock.acquire(blocking=False):
        add_log("Scanare deja in curs.")
        return
    try:
        add_log("Scanez folderul Drive...")
        files = list_pdf_files(FOLDER_ID)
        add_log(f"{len(files)} PDF-uri gasite in total")
        with state_lock:
            known = {fid: f for fid, f in state["files"].items()}
        # fisiere noi sau modificate (re-urcate in Drive) + cele cu eroare
        todo = [f for f in files
                if f["id"] not in known
                or known[f["id"]].get("modifiedTime") != f.get("modifiedTime")
                or known[f["id"]].get("status") == "error"]
        with state_lock:
            for f in todo:
                prev = state["files"].get(f["id"], {})
                state["files"][f["id"]] = {
                    **prev, "id": f["id"], "name": f["name"], "status": "new",
                    "modifiedTime": f.get("modifiedTime"),
                    "interventii": prev.get("interventii", []), "error": None,
                    "avertismente": prev.get("avertismente", []),
                }
            state["last_scan"] = datetime.now().isoformat()
        if todo:
            process_queue_background(todo)
        else:
            add_log("Nicio factura noua.")
    except Exception as e:
        add_log(f"Eroare scanare: {e}")
    finally:
        scan_lock.release()


def monitor_loop():
    while True:
        time.sleep(POLL_INTERVAL)
        with state_lock:
            monitoring = state["monitoring"]
        if monitoring:
            scan_folder()


_bg_started = False


def start_background():
    """Porneste monitorizarea si scanarea initiala — si sub gunicorn, nu doar la `python app.py`."""
    global _bg_started
    if _bg_started:
        return
    _bg_started = True
    threading.Thread(target=monitor_loop, daemon=True).start()
    add_log("Aplicatie pornita (parser direct, AI doar ca rezerva).")
    if FOLDER_ID:
        threading.Thread(target=scan_folder, daemon=True).start()


# ---------------------------------------------------------------- agregari
def aggregate_props(interventii):
    props = {}
    for iv in interventii:
        name = (iv.get("proprietate") or "Necunoscut").strip()
        key = re.sub(r"\s+", " ", name).casefold()
        p = props.setdefault(key, {"name": name, "interventii": 0, "forfait": 0, "ore": 0.0,
                                   "total": 0.0, "clients": set(), "ultima": ""})
        p["interventii"] += 1
        if iv.get("tip_facturare") == "heure":
            p["ore"] += float(iv.get("ore") or iv.get("qty") or 0)
        else:
            p["forfait"] += 1
        p["total"] += float(iv.get("suma") or 0)
        if iv.get("client"):
            p["clients"].add(iv["client"])
        if (iv.get("data_iso") or "") > p["ultima"]:
            p["ultima"] = iv["data_iso"]
    out = list(props.values())
    max_c = max((p["interventii"] for p in out), default=1) or 1
    for p in out:
        p["client"] = ", ".join(sorted(p.pop("clients")))
        p["ore"] = round(p["ore"], 2)
        p["total"] = round(p["total"], 2)
        p["medie"] = round(p["total"] / p["interventii"], 2) if p["interventii"] else 0
        p["pct"] = round(p["interventii"] / max_c * 100)
        p["perf"] = ("Performant" if p["interventii"] >= max_c * 0.7
                     else "Mediu" if p["interventii"] >= max_c * 0.4 else "Rar")
    return sorted(out, key=lambda p: (p["interventii"], p["total"]), reverse=True)


def build_excel():
    with state_lock:
        interventii = list(state["interventii"])
        files = [dict(f) for f in state["files"].values()]
    props = aggregate_props(interventii)

    wb = openpyxl.Workbook()
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    side = Side(style="thin", color="D1D5DB")
    thin = Border(left=side, right=side, top=side, bottom=side)

    def sheet(ws, headers, widths):
        ws.append(headers)
        for col in range(1, len(headers) + 1):
            c = ws.cell(1, col)
            c.fill, c.font, c.border = header_fill, header_font, thin
            c.alignment = Alignment(horizontal="center")
        for i, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        ws.freeze_panes = "A2"

    ws1 = wb.active
    ws1.title = "Clasament proprietati"
    sheet(ws1, ["Rang", "Proprietate", "Client", "Interventii", "din care forfait", "Ore (la ora)",
                "Venit total (EUR)", "Medie / interventie", "Ultima interventie", "Performanta"],
          [6, 28, 26, 12, 14, 12, 16, 18, 16, 13])
    for i, p in enumerate(props, 1):
        ultima = datetime.fromisoformat(p["ultima"]).strftime("%d/%m/%Y") if p["ultima"] else ""
        ws1.append([i, p["name"], p["client"], p["interventii"], p["forfait"], p["ore"] or None,
                    p["total"], p["medie"], ultima, p["perf"]])
        fill = {"Performant": "D1FAE5", "Rar": "FEE2E2"}.get(p["perf"])
        for col in range(1, 11):
            c = ws1.cell(i + 1, col)
            c.border = thin
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)

    ws2 = wb.create_sheet("Detalii interventii")
    sheet(ws2, ["Proprietate", "Client", "Data", "Tip serviciu", "Facturare", "Ore",
                "Pret unitar", "Suma (EUR)", "Factura", "Fisier"],
          [26, 26, 12, 14, 11, 8, 11, 12, 13, 30])
    for iv in sorted(interventii, key=lambda x: (x.get("data_iso") or "", x.get("proprietate") or "")):
        d = datetime.fromisoformat(iv["data_iso"]) if iv.get("data_iso") else None
        ws2.append([iv.get("proprietate"), iv.get("client"), d, iv.get("tip_serviciu"),
                    iv.get("tip_facturare"), iv.get("ore"), iv.get("pret_unitar"),
                    round(float(iv.get("suma") or 0), 2), iv.get("factura"), iv.get("_source")])
        r = ws2.max_row
        ws2.cell(r, 3).number_format = "DD/MM/YYYY"
        for col in range(1, 11):
            ws2.cell(r, col).border = thin

    ws3 = wb.create_sheet("Facturi")
    sheet(ws3, ["Fisier", "Numar", "Client", "Data emiterii", "Total (EUR)", "Interventii",
                "Status", "Metoda", "Avertismente / eroare"],
          [32, 13, 26, 13, 12, 11, 11, 9, 70])
    for f in sorted(files, key=lambda f: f.get("name") or ""):
        ws3.append([f.get("name"), f.get("numar"), f.get("client"), f.get("data_emitere"),
                    f.get("total"), len(f.get("interventii") or []), f.get("status"), f.get("metoda"),
                    f.get("error") or " | ".join(f.get("avertismente") or [])])
        r = ws3.max_row
        for col in range(1, 10):
            ws3.cell(r, col).border = thin
        if f.get("status") in ("warning", "error", "duplicate"):
            ws3.cell(r, 7).fill = PatternFill("solid", fgColor="FEF3C7" if f["status"] != "error" else "FEE2E2")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- rute
def _public_file(f):
    return {k: v for k, v in f.items() if not k.startswith("_")}


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/state")
def api_state():
    with state_lock:
        interventii = list(state["interventii"])
        files = [_public_file(f) for f in state["files"].values()]
        s = {
            "files": files,
            "last_scan": state["last_scan"],
            "monitoring": state["monitoring"],
            "log": state["log"][-50:],
        }
    s["metrics"] = {
        "files": len(files),
        "analyzed": sum(1 for f in files if f["status"] in ("done", "warning", "duplicate")),
        "warnings": sum(1 for f in files if f["status"] in ("warning", "error")),
        "interventii": len(interventii),
        "revenue": round(sum(float(iv.get("suma") or 0) for iv in interventii), 2),
    }
    s["props"] = aggregate_props(interventii)
    return jsonify(s)


@app.route("/api/factura/<file_id>")
def api_factura(file_id):
    with state_lock:
        f = state["files"].get(file_id)
        return (jsonify(_public_file(f)), 200) if f else (jsonify({"error": "not found"}), 404)


@app.route("/api/scan", methods=["POST"])
def api_scan():
    threading.Thread(target=scan_folder, daemon=True).start()
    return jsonify({"ok": True})


@app.route("/api/monitor", methods=["POST"])
def api_monitor():
    data = request.json or {}
    with state_lock:
        state["monitoring"] = bool(data.get("active", False))
        active = state["monitoring"]
    add_log("Monitorizare " + ("activata" if active else "oprita"))
    return jsonify({"monitoring": active})


@app.route("/api/export")
def api_export():
    buf = build_excel()
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"analiza-curatenie-{datetime.now():%Y-%m-%d}.xlsx",
    )


@app.route("/healthz")
def healthz():
    return "ok"


@app.route("/api/test")
def api_test():
    return jsonify({
        "ai_key": bool(LLM_API_KEY),
        "ai_model": LLM_MODEL,
        "folder_id": bool(FOLDER_ID),
        "files_count": len(state["files"]),
        "monitor_thread": _bg_started,
    })


@app.route("/api/test-ai")
def api_test_ai():
    req = urllib.request.Request(f"{LLM_BASE_URL}/models",
                                 headers={"Authorization": f"Bearer {LLM_API_KEY}"})
    try:
        with urllib.request.urlopen(req, timeout=20) as r:
            ids = [m.get("id") for m in json.loads(r.read()).get("data", [])]
            return jsonify({"ok": True, "model_configurat": LLM_MODEL,
                            "model_disponibil": LLM_MODEL in ids, "modele": ids})
    except urllib.error.HTTPError as e:
        return jsonify({"ok": False, "status": e.code, "body": e.read().decode()[:500]}), e.code


@app.route("/api/reset", methods=["POST"])
def api_reset():
    with state_lock:
        state.update({"files": {}, "interventii": [], "last_scan": None,
                      "monitoring": False, "log": []})
    add_log("Reset complet.")
    return jsonify({"ok": True})


# pornim firele de fundal la import (gunicorn importa modulul, nu ruleaza __main__)
if os.environ.get("DISABLE_BACKGROUND") != "1":
    start_background()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port, debug=False)
